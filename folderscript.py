import random
import string
from pathlib import Path
from reportlab.pdfgen import canvas
from docx import Document


def generate_contextual_text(category, filename, word_count=550):
    """Generates varied content based on the file category."""
    themes = {
        'school': ["algebra", "history", "biology", "lecture", "assignment", "textbook", "semester"],
        'office': ["quarterly", "meeting", "stakeholders", "budget", "compliance", "strategy", "deliverables"],
        'story': ["character", "journey", "mystery", "chapter", "plot", "protagonist", "climax"],
        'legal': ["contract", "agreement", "clause", "liability", "terms", "policy", "regulation"]
    }

    # Pick the right vocabulary based on category
    vocab = themes.get(category, ["data", "file", "information", "general"])
    keywords = filename.replace('_', ' ').replace('.', ' ').split()

    body_text = []
    while len(body_text) < word_count:
        # Mix vocabulary, filename keywords, and standard filler
        sentence = [
            random.choice(["Furthermore,", "Interestingly,", "In relation to", "Note that"]),
            random.choice(vocab),
            "is closely tied to",
            random.choice(keywords),
            "because the overall",
            random.choice(vocab),
            "requires a deep dive into the",
            random.choice(vocab),
            "framework."
        ]
        body_text.extend(sentence)

    return " ".join(body_text[:word_count])


def simulate_folder(directory_path, n_files=50):
    path = Path(directory_path)
    path.mkdir(parents=True, exist_ok=True)

    # Expanded variety
    categories = {
        'school': ['biology_notes', 'math_homework', 'history_essay', 'physics_lab'],
        'office': ['annual_report', 'marketing_strategy', 'client_invoice', 'internal_audit'],
        'story': ['fantasy_novel', 'short_story', 'mystery_draft', 'hero_journey'],
        'legal': ['service_agreement', 'privacy_policy', 'liability_waiver', 'nda_draft']
    }

    extensions = ['.txt', '.pdf', '.docx']

    print(f"Generating {n_files} files with deep content in: {path.absolute()}")

    for _ in range(n_files):
        # Choose a random category and a random prefix from it
        cat_key = random.choice(list(categories.keys()))
        prefix = random.choice(categories[cat_key])

        ext = random.choice(extensions)
        rand_id = ''.join(random.choices(string.digits, k=3))
        filename = f"{prefix}_{rand_id}{ext}"

        file_path = path / filename
        content = generate_contextual_text(cat_key, filename)

        if ext == ".txt":
            file_path.write_text(content)

        elif ext == ".docx":
            doc = Document()
            doc.add_heading(filename.replace('_', ' ').title(), 0)
            doc.add_paragraph(content)
            doc.save(file_path)

        elif ext == ".pdf":
            c = canvas.Canvas(str(file_path))
            t = c.beginText(50, 800)
            t.setFont("Helvetica", 10)

            # Simple text wrapping for PDF
            chars_per_line = 95
            lines = [content[i:i + chars_per_line] for i in range(0, len(content), chars_per_line)]

            for line in lines:
                if t.getY() < 50:  # Start new page if at bottom
                    c.drawText(t)
                    c.showPage()
                    t = c.beginText(50, 800)
                    t.setFont("Helvetica", 10)
                t.textLine(line)

            c.drawText(t)
            c.save()

    print("Simulation complete. Files are ready for AI training.")


# --- EXECUTION ---
test_folder = "/home/dan/Documents"
simulate_folder(test_folder, n_files=200)